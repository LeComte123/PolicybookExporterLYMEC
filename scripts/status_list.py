"""The status list resolution_status_list.xlsx: the Bureau's record of every Policy Book number.

Tab "Resolutions", one row per number, columns found by their header (Number, Title, Status,
Replaced by (number), Archived at (congress id), ...). Status: "in force", "archived" or "not used".

The list confirms archiving: a motion text saying it replaces or archives a resolution only takes
effect when the list marks that resolution as archived (import_openslides.py), and every resolution
marked archived in the list is archived (validate.py --fix). Rows for new resolutions are added by
validate.py --fix. The migration of the old Policy Book uses the same rules.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from rapidfuzz import fuzz

import pbook

FILE = "resolution_status_list.xlsx"
SHEET = "Resolutions"
TITLE_MATCH = 85  # rapidfuzz ratio of two normalised titles that are the same resolution
COLUMNS = {  # field: header in the sheet (start of the header text)
    "number": "Number",
    "title": "Title",
    "status": "Status",
    "replaced_by": "Replaced by",
    "archived_at": "Archived at",
    "adopted_at": "Adopted at",
}


class StatusListError(Exception):
    pass


@dataclass
class Row:
    line: int  # row number in the sheet, for messages
    number: str
    title: str
    status: str  # lower case: "in force", "archived", "not used" or ""
    replaced_by: str
    archived_at: str


def path(root: Path) -> Path:
    return root / FILE


def as_number(value) -> str:
    """'6.10', or 6.1 when Excel turned the cell into a number."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{value:.2f}"
    return str(value or "").strip()


def _columns(sheet) -> dict[str, int]:
    headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
    columns = {}
    for field, header in COLUMNS.items():
        column = next((col for text, col in headers.items() if text.startswith(header)), None)
        if column is None and field in ("number", "status"):
            raise StatusListError(f"{FILE}: the tab '{SHEET}' has no column '{header}'. Restore the column header in row 1.")
        columns[field] = column
    return columns


def load(root: Path) -> list[Row] | None:
    """The rows of the status list, or None if there is no list."""
    if not path(root).exists():
        return None
    try:
        workbook = load_workbook(path(root), read_only=True, data_only=True)
    except Exception as exc:  # a damaged or non-Excel file
        raise StatusListError(f"{FILE} cannot be opened ({exc}). Upload the file again, saved as Excel workbook (.xlsx).") from None
    if SHEET not in workbook.sheetnames:
        raise StatusListError(f"{FILE} has no tab '{SHEET}'. Rename the tab back to '{SHEET}'.")
    sheet = workbook[SHEET]
    columns = _columns(sheet)
    rows = []
    for line, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        def get(field):
            column = columns.get(field)
            return values[column - 1] if column and column <= len(values) else None

        number = as_number(get("number"))
        if not number:
            continue
        rows.append(Row(
            line=line,
            number=number,
            title=str(get("title") or "").strip(),
            status=str(get("status") or "").strip().lower(),
            replaced_by=as_number(get("replaced_by")),
            archived_at=str(get("archived_at") or "").strip(),
        ))  # fmt: skip
    workbook.close()
    return rows


def normalise(title: str) -> str:
    return " ".join(re.sub(r"[^\w\s]", " ", title.casefold()).split())


def titles_match(a: str, b: str) -> bool:
    return fuzz.ratio(normalise(a), normalise(b)) >= TITLE_MATCH


def archived_row(rows: list[Row] | None, number: str, title: str | None = None) -> Row | None:
    """The row marking this number as archived; if the row has a title, it must match `title`."""
    for row in rows or []:
        if row.number == number and row.status == "archived" and (not row.title or not title or titles_match(row.title, title)):
            return row
    return None


def add_rows(root: Path, resolutions: list[pbook.Resolution]) -> None:
    """Write rows for these resolutions (status 'in force') into the first empty rows of the list."""
    workbook = load_workbook(path(root))
    sheet = workbook[SHEET]
    columns = _columns(sheet)
    line = 2
    for res in resolutions:
        while sheet.cell(row=line, column=columns["number"]).value not in (None, ""):
            line += 1
        values = {"number": res.number, "title": res.meta.get("title"), "status": "in force", "adopted_at": res.meta.get("congress")}
        for field, value in values.items():
            if columns.get(field):
                sheet.cell(row=line, column=columns[field]).value = str(value) if value is not None else None
    workbook.save(path(root))
